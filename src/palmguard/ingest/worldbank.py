"""World Bank Commodity Markets (Pink Sheet) monthly xlsx ingestor and parser."""

from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urljoin

import httpx
import openpyxl
import pandas as pd

from palmguard.http import fetch_url
from palmguard.ingest.snapshots import (
    SnapshotValidationError,
    SnapshotWriteResult,
    write_snapshot,
)

WORLDBANK_LANDING_URL = "https://www.worldbank.org/en/research/commodity-markets"
WORLDBANK_SOURCE = "worldbank"
PINK_SHEET_DATASET = "pink_sheet"
PINK_SHEET_FILENAME = "CMO-Historical-Data-Monthly.xlsx"

REQUIRED_COLUMNS: dict[str, str] = {
    "Palm oil": "palm_wb_usd",
    "Crude oil, Brent": "brent_wb_usd",
    "Soybean oil": "soyoil_wb_usd",
}

EXPECTED_UNITS: dict[str, str] = {
    "Palm oil": "($/mt)",
    "Crude oil, Brent": "($/bbl)",
    "Soybean oil": "($/mt)",
}

_MONTH_CODE_RE = re.compile(r"^(\d{4})M(0[1-9]|1[0-2])$")
_HREF_RE = re.compile(
    r"""href\s*=\s*["']([^"']*CMO-Historical-Data-Monthly\.xlsx[^"']*)["']""", re.I
)


@dataclass(frozen=True)
class PinkSheetMetadata:
    """Metadata and commodity descriptions extracted directly from the Pink Sheet workbook."""

    updated_note: str
    columns_found: dict[str, str]
    start_months: dict[str, str]
    end_months: dict[str, str]
    row_counts: dict[str, int]
    descriptions: dict[str, str]


def resolve_monthly_xlsx_url(landing_html: str, *, base_url: str = WORLDBANK_LANDING_URL) -> str:
    """Resolve the current CMO-Historical-Data-Monthly.xlsx link from the landing page HTML."""
    match = _HREF_RE.search(landing_html)
    if not match:
        raise SnapshotValidationError(
            "Could not find CMO-Historical-Data-Monthly.xlsx link on World Bank landing page"
        )
    href = match.group(1).strip()
    return urljoin(base_url, href)


def _clean_header(val: object) -> str:
    if val is None:
        return ""
    return re.sub(r"\s+", " ", str(val)).strip()


def _parse_numeric_cell(val: object, *, series_id: str, month_str: str) -> float | None:
    if val is None:
        return None
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        fval = float(val)
        if math.isnan(fval) or math.isinf(fval):
            return None
        if fval <= 0.0:
            raise SnapshotValidationError(
                f"Non-positive value {fval} for {series_id} at {month_str}"
            )
        return fval

    sval = str(val).strip()
    if not sval or sval in {"..", "...", "—", "-", "N/A", "n.a.", "\ufffd"}:
        return None
    try:
        fval = float(sval)
    except ValueError:
        return None
    if not math.isfinite(fval):
        return None
    if fval <= 0.0:
        raise SnapshotValidationError(f"Non-positive value {fval} for {series_id} at {month_str}")
    return fval


def _extract_descriptions(wb: openpyxl.Workbook) -> dict[str, str]:
    descriptions: dict[str, str] = {}
    if "Description" not in wb.sheetnames:
        return descriptions
    ws = wb["Description"]
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
        if not cells:
            continue
        joined = " | ".join(cells)
        for text in cells:
            lower = text.lower()
            if lower.startswith("palm oil (") or lower.startswith("palm oil,"):
                descriptions["Palm oil"] = text
            elif lower.startswith("crude oil, uk brent") or lower.startswith("crude oil, brent"):
                descriptions["Crude oil, Brent"] = text
            elif lower.startswith("soybean oil,") or lower.startswith("soybean oil ("):
                descriptions["Soybean oil"] = text
        if "replacement series for palm oil" in joined.lower():
            descriptions["Palm oil revision note"] = joined
    return descriptions


def parse_pink_sheet_xlsx(
    payload: bytes,
    *,
    snapshot_id: str = "",
    available_from: str = "",
) -> tuple[pd.DataFrame, PinkSheetMetadata]:
    """Parse Monthly Prices and Description sheets from a World Bank Pink Sheet xlsx payload."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    except Exception as exc:
        raise SnapshotValidationError(f"Invalid xlsx workbook: {exc}") from exc

    try:
        if "Monthly Prices" not in wb.sheetnames:
            raise SnapshotValidationError("Missing 'Monthly Prices' sheet in Pink Sheet workbook")

        ws = wb["Monthly Prices"]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 7:
            raise SnapshotValidationError("Monthly Prices sheet has too few rows")

        updated_note = ""
        header_idx: int | None = None
        for idx, row in enumerate(rows[:15]):
            non_empty = [_clean_header(c) for c in row if c is not None and _clean_header(c)]
            for text in non_empty:
                if text.lower().startswith("updated on"):
                    updated_note = text
            cleaned_row = [_clean_header(c) for c in row]
            if "Palm oil" in cleaned_row and "Crude oil, Brent" in cleaned_row:
                header_idx = idx
                break

        if header_idx is None or header_idx + 1 >= len(rows):
            raise SnapshotValidationError(
                "Required columns ('Palm oil', 'Crude oil, Brent', 'Soybean oil') not found"
            )

        header_row = [_clean_header(c) for c in rows[header_idx]]
        unit_row = [_clean_header(c) for c in rows[header_idx + 1]]

        col_indices: dict[str, int] = {}
        columns_found: dict[str, str] = {}
        for col_name in REQUIRED_COLUMNS:
            if col_name not in header_row:
                raise SnapshotValidationError(f"Missing required column: {col_name}")
            c_idx = header_row.index(col_name)
            unit = unit_row[c_idx] if c_idx < len(unit_row) else ""
            expected_unit = EXPECTED_UNITS[col_name]
            if expected_unit and unit != expected_unit:
                raise SnapshotValidationError(
                    f"Unexpected unit for {col_name}: got {unit!r}, expected {expected_unit!r}"
                )
            col_indices[col_name] = c_idx
            columns_found[col_name] = unit

        records: list[dict[str, object]] = []
        for row in rows[header_idx + 2 :]:
            if not row or row[0] is None:
                continue
            raw_code = str(row[0]).strip()
            m = _MONTH_CODE_RE.match(raw_code)
            if not m:
                continue
            month_str = f"{m.group(1)}-{m.group(2)}-01"
            for col_name, series_id in REQUIRED_COLUMNS.items():
                c_idx = col_indices[col_name]
                cell_val = row[c_idx] if c_idx < len(row) else None
                val = _parse_numeric_cell(cell_val, series_id=series_id, month_str=month_str)
                if val is not None:
                    records.append(
                        {
                            "series_id": series_id,
                            "month": month_str,
                            "value": val,
                            "available_from": available_from,
                            "snapshot_id": snapshot_id,
                            "is_provisional": False,
                        }
                    )

        if not records:
            raise SnapshotValidationError("No monthly price rows parsed from Pink Sheet")

        df = pd.DataFrame(records).sort_values(["series_id", "month"]).reset_index(drop=True)

        # Check for duplicate months within any series
        if df.duplicated(subset=["series_id", "month"]).any():
            raise SnapshotValidationError("Duplicate (series_id, month) rows in Pink Sheet")

        start_months: dict[str, str] = {}
        end_months: dict[str, str] = {}
        row_counts: dict[str, int] = {}
        for col_name, series_id in REQUIRED_COLUMNS.items():
            sub = df[df["series_id"] == series_id]
            if sub.empty:
                raise SnapshotValidationError(
                    f"Zero valid observations for {series_id} ({col_name})"
                )
            start_months[series_id] = str(sub["month"].iloc[0])
            end_months[series_id] = str(sub["month"].iloc[-1])
            row_counts[series_id] = len(sub)

        descriptions = _extract_descriptions(wb)
        meta = PinkSheetMetadata(
            updated_note=updated_note,
            columns_found=columns_found,
            start_months=start_months,
            end_months=end_months,
            row_counts=row_counts,
            descriptions=descriptions,
        )
        return df, meta
    finally:
        wb.close()


def validate_pink_sheet_xlsx(payload: bytes) -> int:
    """Validate raw Pink Sheet xlsx bytes and return total parsed row count."""
    df, _ = parse_pink_sheet_xlsx(payload)
    return len(df)


def ingest_worldbank_pink_sheet(
    raw_root: Path,
    client: httpx.Client,
    *,
    landing_url: str = WORLDBANK_LANDING_URL,
    retrieved_at_utc: str | None = None,
) -> tuple[SnapshotWriteResult, pd.DataFrame, PinkSheetMetadata]:
    """Resolve the monthly Pink Sheet xlsx from the landing page, validate, and store snapshot."""
    landing_res = fetch_url(client, landing_url)
    xlsx_url = resolve_monthly_xlsx_url(
        landing_res.content.decode("utf-8", errors="replace"),
        base_url=landing_res.resolved_url,
    )
    xlsx_res = fetch_url(client, xlsx_url)
    effective_ts = retrieved_at_utc or xlsx_res.retrieved_at_utc

    df, meta = parse_pink_sheet_xlsx(xlsx_res.content)
    write_res = write_snapshot(
        raw_root,
        source=WORLDBANK_SOURCE,
        dataset=PINK_SHEET_DATASET,
        payload=xlsx_res.content,
        filename=PINK_SHEET_FILENAME,
        request_url=xlsx_url,
        resolved_url=xlsx_res.resolved_url,
        http_status=xlsx_res.http_status,
        retrieved_at_utc=effective_ts,
        row_count=len(df),
    )

    avail_date = write_res.manifest.retrieved_at_utc[:10]
    df = df.assign(
        available_from=avail_date,
        snapshot_id=write_res.manifest.snapshot_id,
    )
    return write_res, df, meta
